"""Làm lại module Tài sản 05/10/2026 — Thôi dùng, Bỏ thôi dùng, chặn Xoá, trần 36 tháng của công cụ
dụng cụ, nạp máy đã khấu hao hết, nhập tài sản đang dùng từ Excel.

Spec: `docs/superpowers/specs/2026-10-05-tai-san-lam-lai-design.md`.
"""
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db import Base
import app.models  # noqa: F401
from app.models.department import Department
from app.models.tai_san import (
    BD_THOI_DUNG,
    LOAI_CCDC,
    LOAI_TSCD,
    TT_DA_GIAM,
    TT_DANG_DUNG,
    TaiSan,
    TaiSanBienDong,
)
from app.repositories.tai_san_repo import TaiSanRepository
from app.services.tai_san.bang_thang import bang_thang
from app.services.tai_san.nhap_excel import COT, ExcelSaiMan, nhap_dang_dung, tao_mau
from app.services.tai_san.service import (
    TaiSanDaCoChungTu,
    TaiSanService,
    TaiSanValidationError,
)


def _svc():
    eng = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(eng)
    db = sessionmaker(bind=eng)()
    return db, TaiSanService(TaiSanRepository(db))


def _komori(svc):
    """3,3 tỷ ÷ 120 tháng = 27.500.000/tháng, dùng từ 10/03/2026."""
    return svc.ghi_tang(dict(
        ten="May in Komori", loai=LOAI_TSCD, so_thang=120, ngay_su_dung=date(2026, 3, 10),
        chi_phi=[{"dien_giai": "Gia mua", "so_tien": 3_300_000_000}],
    ))


# --- Thôi dùng ---------------------------------------------------------------------------


def test_thoi_dung_giua_thang_chia_theo_ngay_roi_ngung():
    db, svc = _svc()
    t = _komori(svc)
    svc.thoi_dung(t.id, ngay=date(2026, 9, 15), kieu="ban", ly_do="Ban cho xuong ban")
    assert t.trang_thai == TT_DA_GIAM and t.ngay_giam == date(2026, 9, 15)
    r = bang_thang(db, 2026, 9)[0]
    assert r["muc_trich"] == 27_500_000 * 15 // 30 and r["con_lai"] == 0
    assert [(s["loai"], s["nhan"]) for s in r["su_kien"]] == [
        ("thoi_dung", "Thôi dùng từ 15/09 (Bán)"),
    ]
    assert r["dien_giai"] == (
        "Bán ngày 15/09/2026: ngừng khấu hao từ ngày này. Ban cho xuong ban"
    )
    assert bang_thang(db, 2026, 10) == []
    bd = db.query(TaiSanBienDong).one()
    assert bd.loai == BD_THOI_DUNG and bd.kieu_thoi_dung == "ban"


def test_thoi_dung_chan_ly_do_ngay_va_lan_hai():
    db, svc = _svc()
    t = _komori(svc)
    with pytest.raises(TaiSanValidationError, match="Chọn lý do"):
        svc.thoi_dung(t.id, ngay=date(2026, 9, 1), kieu="cho_muon")
    with pytest.raises(TaiSanValidationError, match="trước ngày bắt đầu dùng"):
        svc.thoi_dung(t.id, ngay=date(2026, 3, 1), kieu="hong")
    bp = Department(name="To Be", code="PB901")
    db.add(bp)
    db.commit()
    svc.dieu_chuyen(t.id, ngay=date(2026, 6, 1), bo_phan_moi_id=bp.id)
    with pytest.raises(TaiSanValidationError, match="sau lần thay đổi gần nhất"):
        svc.thoi_dung(t.id, ngay=date(2026, 5, 1), kieu="hong")
    svc.thoi_dung(t.id, ngay=date(2026, 6, 1), kieu="hong")
    with pytest.raises(TaiSanValidationError, match="đã thôi dùng"):
        svc.thoi_dung(t.id, ngay=date(2026, 7, 1), kieu="mat")


def test_da_thoi_dung_thi_khong_chuyen_khong_sua_lon():
    db, svc = _svc()
    t = _komori(svc)
    svc.thoi_dung(t.id, ngay=date(2026, 9, 15), kieu="thanh_ly")
    with pytest.raises(TaiSanValidationError, match="đã thôi dùng"):
        svc.nang_cap(t.id, ngay=date(2026, 10, 1), so_tien=1_000_000, so_thang_con_lai=50)
    with pytest.raises(TaiSanValidationError, match="đã thôi dùng"):
        svc.dieu_chuyen(t.id, ngay=date(2026, 10, 1), bo_phan_moi_id=1)


def test_bo_thoi_dung_tra_lai_lich_nhu_cu():
    db, svc = _svc()
    t = _komori(svc)
    truoc = [(d.nam, d.thang, d.muc_trich) for d in svc.lich(t)]
    svc.thoi_dung(t.id, ngay=date(2026, 9, 15), kieu="mat")
    svc.bo_thoi_dung(t.id)
    assert t.trang_thai == TT_DANG_DUNG and t.ngay_giam is None
    assert db.query(TaiSanBienDong).count() == 0
    assert [(d.nam, d.thang, d.muc_trich) for d in svc.lich(t)] == truoc
    with pytest.raises(TaiSanValidationError, match="đang dùng"):
        svc.bo_thoi_dung(t.id)
    # Bỏ xong lại xoá được (không còn lịch sử)
    svc.xoa(t.id)


def test_xoa_tai_san_da_thoi_dung_bi_chan():
    db, svc = _svc()
    t = _komori(svc)
    svc.thoi_dung(t.id, ngay=date(2026, 9, 15), kieu="hong")
    with pytest.raises(TaiSanDaCoChungTu, match="không xoá được"):
        svc.xoa(t.id)


# --- Trần 36 tháng của công cụ dụng cụ ------------------------------------------------------


def test_cong_cu_dung_cu_toi_da_36_thang():
    db, svc = _svc()
    lo = dict(ten="Pallet go", loai=LOAI_CCDC, so_luong=10, don_gia=2_000_000,
              ngay_su_dung=date(2026, 7, 1))
    with pytest.raises(TaiSanValidationError, match="tối đa 36 tháng"):
        svc.ghi_tang({**lo, "so_thang": 48})
    t = svc.ghi_tang({**lo, "so_thang": 36})
    with pytest.raises(TaiSanValidationError, match="tối đa 36 tháng"):
        svc.sua(t.id, {"so_thang": 40})
    # Tài sản cố định thì không có trần này
    svc.ghi_tang(dict(ten="May cat", loai=LOAI_TSCD, so_thang=180, don_gia=500_000_000,
                      ngay_su_dung=date(2026, 7, 1)))


# --- Máy đã khấu hao hết mà vẫn chạy ------------------------------------------------------


def test_nap_may_da_khau_hao_het_khong_trich_thang_nao():
    db, svc = _svc()
    t = svc.nap_dau_ky(dict(
        ten="May gap cu", loai=LOAI_TSCD, so_thang=84, don_gia=200_000_000,
        ngay_su_dung=date(2015, 1, 1), moc_tu_ngay=date(2026, 10, 1),
        thang_da_trich_dau_ky=84, hao_mon_dau_ky=200_000_000,
    ))
    assert svc.lich(t) == []
    assert svc.hao_mon_den(t, 2026, 12) == 200_000_000
    assert bang_thang(db, 2026, 10) == []


# --- Nhập Excel ---------------------------------------------------------------------------


def _file(*dong):
    wb = Workbook()
    ws = wb.active
    ws.append([t for _, t, _ in COT])
    for d in dong:
        ws.append(list(d))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_file_mau_co_du_cot_va_huong_dan():
    wb = load_workbook(BytesIO(tao_mau()))
    tieu_de = [c.value for c in wb.worksheets[0][1]]
    assert tieu_de[0] == "Tên tài sản *" and len(tieu_de) == len(COT)
    assert "Hướng dẫn" in wb.sheetnames


def test_nhap_excel_tu_dien_va_xem_truoc_khong_ghi():
    db, svc = _svc()
    db.add(Department(name="To In", code="PB902"))
    db.commit()
    du_lieu = _file(
        # Tên, Loại, Giá một cái, SL, Tháng, Bắt đầu, Đã KH tháng, Đã KH đồng, Bộ phận, Tính từ
        ("May in Heidelberg", None, 1_200_000_000, None, 120, "01/10/2020", None, None, "to in",
         "10/2026"),
        ("Xe day hang", None, 3_000_000, 4, 24, date(2026, 1, 15), None, None, None, "10/2026"),
        ("May gap cu", "Tài sản cố định", 200_000_000, None, 84, "01/01/2015", 84, 200_000_000,
         None, "10/2026"),
    )
    kq = nhap_dang_dung(db, svc, du_lieu, user_id=None, ghi=False)
    assert kq.loi == [] and kq.tong_dong == 3 and kq.tao_moi == 3 and not kq.da_ghi
    assert db.query(TaiSan).count() == 0

    kq = nhap_dang_dung(db, svc, du_lieu, user_id=None, ghi=True)
    assert kq.da_ghi
    heidel, xe, gap = db.query(TaiSan).order_by(TaiSan.id).all()
    # 10/2020 → 10/2026 = 72 tháng; 1,2 tỷ × 72 ÷ 120
    assert heidel.loai == LOAI_TSCD and heidel.thang_da_trich_dau_ky == 72
    assert heidel.hao_mon_dau_ky == 720_000_000 and heidel.bo_phan_id is not None
    assert heidel.ma.startswith("TS-") and gap.ma != heidel.ma
    # 3 triệu một cái < 30 triệu ⇒ công cụ dụng cụ; tổng = 4 cái
    assert xe.loai == LOAI_CCDC and xe.nguyen_gia == 12_000_000 and xe.ma.startswith("CC-")
    assert xe.thang_da_trich_dau_ky == 9 and xe.hao_mon_dau_ky == 12_000_000 * 9 // 24
    assert svc.lich(gap) == []


def test_nhap_excel_mot_dong_sai_la_khong_ghi_gi():
    db, svc = _svc()
    du_lieu = _file(
        ("May in Heidelberg", None, 1_200_000_000, None, 120, "01/10/2020"),
        ("Pallet", "Công cụ dụng cụ", 500_000, 10, 48, "01/01/2026"),
        ("May cat", None, 400_000_000, 2, 120, "01/01/2024"),
        ("", None, None, None, None, None),
        ("May dan", None, 90_000_000, None, 96, "ngay mai", None, None, "Khong co"),
    )
    kq = nhap_dang_dung(db, svc, du_lieu, user_id=None, ghi=True, hom_nay=date(2026, 10, 5))
    assert not kq.da_ghi and db.query(TaiSan).count() == 0
    assert kq.tong_dong == 4
    assert [(d, c) for d, c, _ in kq.loi] == [
        (3, ""), (4, "Số lượng"), (6, "Bắt đầu dùng từ ngày"),
    ]
    assert "tối đa 36 tháng" in kq.loi[0][2]


def test_nhap_excel_sai_mau():
    db, svc = _svc()
    wb = Workbook()
    wb.active.append(["Ten", "Gia"])
    out = BytesIO()
    wb.save(out)
    with pytest.raises(ExcelSaiMan, match="thiếu cột"):
        nhap_dang_dung(db, svc, out.getvalue(), user_id=None, ghi=False)
    with pytest.raises(ExcelSaiMan):
        nhap_dang_dung(db, svc, b"khong phai xlsx", user_id=None, ghi=False)


# --- HTTP ---------------------------------------------------------------------------------


def _token(client, seed_credentials):
    r = client.post("/api/auth/login", json=seed_credentials)
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def test_api_thoi_dung_roi_bo(client, seed_credentials):
    h = _token(client, seed_credentials)
    r = client.post("/api/tai-san", headers=h, json={
        "ten": "May in Komori", "loai": "tscd", "so_thang": 120, "ngay_su_dung": "2026-03-10",
        "so_luong": 1, "don_gia": 3300000000,
    })
    assert r.status_code == 201, r.text
    tid = r.json()["id"]
    r = client.post(f"/api/tai-san/{tid}/thoi-dung", headers=h,
                    json={"ngay": "2026-09-15", "kieu": "ban"})
    assert r.status_code == 201, r.text
    assert r.json()["kieu_thoi_dung"] == "ban"
    ds = client.get("/api/tai-san?trang_thai=da_giam", headers=h).json()
    assert [x["id"] for x in ds["items"]] == [tid] and ds["items"][0]["con_lai"] == 0
    assert client.get("/api/tai-san?trang_thai=dang_dung", headers=h).json()["total"] == 0
    assert client.delete(f"/api/tai-san/{tid}", headers=h).status_code == 409
    r = client.delete(f"/api/tai-san/{tid}/thoi-dung", headers=h)
    assert r.status_code == 200 and r.json()["trang_thai"] == "dang_dung"
    assert client.delete(f"/api/tai-san/{tid}/thoi-dung", headers=h).status_code == 422


def test_api_nhap_excel_va_file_mau(client, seed_credentials):
    h = _token(client, seed_credentials)
    r = client.get("/api/tai-san/mau-excel", headers=h)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    du_lieu = _file(("May in Heidelberg", None, 1_200_000_000, None, 120, "01/10/2020"))
    tep = {"file": ("ts.xlsx", du_lieu, "application/octet-stream")}
    r = client.post("/api/tai-san/import-excel?mode=preview", headers=h, files=tep)
    assert r.status_code == 200, r.text
    assert r.json()["hop_le"] and r.json()["tao_moi"] == 1 and not r.json()["da_ghi"]
    assert client.get("/api/tai-san", headers=h).json()["total"] == 0
    r = client.post("/api/tai-san/import-excel?mode=commit", headers=h, files=tep)
    assert r.json()["da_ghi"]
    assert client.get("/api/tai-san", headers=h).json()["total"] == 1


def test_api_tien_sua_chua_lon(client, seed_credentials):
    h = _token(client, seed_credentials)
    tid = client.post("/api/tai-san", headers=h, json={
        "ten": "May in", "loai": "tscd", "so_thang": 120, "ngay_su_dung": "2026-03-10",
        "so_luong": 1, "don_gia": 3300000000,
    }).json()["id"]
    r = client.post(f"/api/tai-san/{tid}/bien-dong", headers=h, json={
        "loai": "nang_cap", "ngay": "2026-04-15", "so_tien": 100000000, "so_thang_con_lai": 118,
    })
    assert r.status_code == 201, r.text
    row = client.get(f"/api/tai-san/{tid}", headers=h).json()
    assert row["nguyen_gia"] == 3400000000 and row["tien_sua_chua_lon"] == 100000000
