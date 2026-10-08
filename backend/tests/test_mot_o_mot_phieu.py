"""Một ô một phiếu (spec `docs/spec-ke-hoach-vat-tu-mot-o-mot-phieu.md`, 07/10/2026).

Ô = mặt hàng (mã + khổ) × lệnh/bài. Ba thứ phải đúng, sai là mua trùng hoặc lệnh mất hàng của mình:

* ô đã có phiếu còn chạy thì KHÔNG đề nghị mua lần hai (cả ở Kế hoạch vật tư lẫn lúc Lưu form);
* đơn mua lập từ yêu cầu của ô thì phần đặt cho lệnh chỉ lệnh đó dùng — lệnh khác hạn sớm hơn
  cũng không lấy được; phần đặt dư là hàng chung;
* giữ chỗ lấy phần đặt cho mình trước.
"""
from __future__ import annotations

from datetime import datetime, timedelta, timezone

import pytest

from app.models.purchase import (
    DPR_OPEN,
    PR_PURCHASED,
    PR_REJECTED,
    SOURCE_SAN_XUAT,
    DepartmentPurchaseRequest,
    DepartmentPurchaseRequestLine,
    PurchaseRequest,
    PurchaseRequestLine,
    Supplier,
    SupplierItem,
    YeuCauMuaNguonLenh,
)
from app.models.vat_tu_giu_cho import NGUON_DANG_VE
from app.services.giu_cho_service import GiuChoService
from app.services.ke_hoach_vat_tu_service import KeHoachVatTuValidationError
from app.services.mach_mua import BUOC_DAT, BUOC_MOI, BUOC_TRA
from app.services.purchase_service import PurchaseService
from tests.test_ke_hoach_vat_tu import (  # noqa: F401 — fixture dùng lại
    HOM_NAY,
    KHO,
    MAI,
    _admin_token,
    _buoc_dau,
    _dv_dong,
    _giay,
    _kho_dong,
    _lenh,
    _nhom,
    customer,
    db,
    svc,
)


def _yc(db, hang, cho: list[tuple], *, so_luong=None):
    """Yêu cầu mua MỘT dòng, liên kết các ô `cho = [(lệnh, số đề nghị | None)]`."""
    yc = DepartmentPurchaseRequest(
        code=f"YCMH-{db.query(DepartmentPurchaseRequest).count() + 1}",
        status=DPR_OPEN, source_type=SOURCE_SAN_XUAT, purpose="Thiếu giấy", needed_date=MAI,
    )
    db.add(yc)
    db.flush()
    ln = DepartmentPurchaseRequestLine(
        department_request_id=yc.id, item_name="Giấy", hang_loai=hang[0], hang_id=hang[1],
        unit=_dv_dong(hang, None),
        quantity=so_luong if so_luong is not None else sum(s or 0 for _, s in cho),
        kho_rong=_kho_dong(hang, KHO)[0], kho_dai=_kho_dong(hang, KHO)[1],
    )
    db.add(ln)
    for lenh, sl in cho:
        db.add(YeuCauMuaNguonLenh(
            department_request_id=yc.id, hang_loai=hang[0], hang_id=hang[1],
            kho_rong=_kho_dong(hang, KHO)[0], kho_dai=_kho_dong(hang, KHO)[1],
            lsx_id=lenh.id, so_luong=sl,
        ))
    db.commit()
    return yc, ln


def _don_tu(db, yc_line, so_luong, *, ngay_ve=MAI, status=PR_PURCHASED):
    """Đơn mua lập từ dòng yêu cầu — đúng đường thu mua đi (`department_request_line_id`)."""
    p = PurchaseRequest(code=f"PMH-{db.query(PurchaseRequest).count() + 1}", status=status,
                        expected_receipt_date=ngay_ve)
    db.add(p)
    db.flush()
    pl = PurchaseRequestLine(
        purchase_request_id=p.id, item_name="Giấy", hang_loai=yc_line.hang_loai,
        hang_id=yc_line.hang_id, kho_rong=yc_line.kho_rong, kho_dai=yc_line.kho_dai,
        unit=yc_line.unit, quantity=so_luong, expected_unit_price=1,
        department_request_line_id=yc_line.id,
    )
    db.add(pl)
    db.commit()
    return p, pl


def _hang(g):
    return ("giay", g.id)


def _khoa(g, lenh, buoc):
    return {"hang_loai": "giay", "hang_id": g.id, "kho_rong": KHO[0], "kho_dai": KHO[1],
            "lsx_id": lenh.id, "bai_ghep_id": None, "buoc_id": buoc.id}


def _dong(bang, g):
    return {d["ma"]: d for d in _nhom(bang, g)["dong"]}


# --- Chặn đề nghị lần hai --------------------------------------------------------


def test_o_da_co_yeu_cau_mo_thi_KHONG_de_nghi_lan_hai(db, svc, customer):
    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500, han=MAI)
    _yc(db, _hang(g), [(a, 500)])

    with pytest.raises(KeHoachVatTuValidationError, match="đã có phiếu YCMH-1"):
        svc.gom_de_nghi([_khoa(g, a, _buoc_dau(db, a))])


def test_phieu_cua_lenh_A_khong_chan_lenh_B(db, svc, customer):
    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500, han=MAI)
    b = _lenh(db, customer, ma="LSX-B", giay_id=g.id, so_to_nguyen=300, han=MAI)
    _yc(db, _hang(g), [(a, 500)])

    gom = svc.gom_de_nghi([_khoa(g, b, _buoc_dau(db, b))])
    assert gom["lines"][0]["quantity"] == pytest.approx(300)
    assert gom["nguon"][0]["so_luong"] == pytest.approx(300), "liên kết phải mang số đề nghị"


def test_huy_dong_yeu_cau_thi_de_nghi_lai_duoc(db, svc, customer):
    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500, han=MAI)
    _, ln = _yc(db, _hang(g), [(a, 500)])
    ln.cancelled_at = datetime.now(timezone.utc)
    db.commit()

    assert svc.gom_de_nghi([_khoa(g, a, _buoc_dau(db, a))])["lines"]


def test_don_bi_tra_van_la_phieu_song_cua_o(db, svc, customer):
    """Đơn bị trả/huỷ mà yêu cầu còn mở: việc của người mua là lập lại đơn, KHÔNG phải đề nghị mới."""
    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500, han=MAI)
    _, ln = _yc(db, _hang(g), [(a, 500)])
    _don_tu(db, ln, 500, status=PR_REJECTED)

    svc.can_doi()
    o = svc.mach["theo_o"][((("giay", g.id) + KHO), a.id, None)]
    assert o["phieu"]["buoc"] == BUOC_TRA
    with pytest.raises(KeHoachVatTuValidationError):
        svc.gom_de_nghi([_khoa(g, a, _buoc_dau(db, a))])


def test_hai_yeu_cau_cu_cung_o_thi_mot_phieu_mot_trung(db, svc, customer):
    """Dữ liệu trước 07/10 có thể có hai yêu cầu cho cùng ô — phiếu là cái đi xa nhất, cái kia trùng."""
    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500, han=MAI)
    _yc(db, _hang(g), [(a, 500)])
    _, ln2 = _yc(db, _hang(g), [(a, 500)])
    _don_tu(db, ln2, 500)

    svc.can_doi()
    o = svc.mach["theo_o"][((("giay", g.id) + KHO), a.id, None)]
    assert o["phieu"]["buoc"] == BUOC_DAT and o["phieu"]["yc_ma"] == "YCMH-2"
    assert [m["yc_ma"] for m in o["trung"]] == ["YCMH-1"]
    assert o["trung"][0]["buoc"] == BUOC_MOI


# --- Phần đặt cho lệnh -----------------------------------------------------------


def test_hang_dat_cho_lenh_A_lenh_B_han_som_hon_KHONG_lay_duoc(db, svc, customer):
    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500,
              han=HOM_NAY + timedelta(days=20))
    b = _lenh(db, customer, ma="LSX-B", giay_id=g.id, so_to_nguyen=500,
              han=HOM_NAY + timedelta(days=5))         # ăn tồn TRƯỚC A
    _, ln = _yc(db, _hang(g), [(a, 500)])
    _don_tu(db, ln, 500)

    dong = _dong(svc.can_doi(), g)
    assert dong["LSX-A"]["thieu"] == pytest.approx(0)
    assert dong["LSX-A"]["trang_thai"] != "xanh", "hàng chưa về thì không phải xanh"
    assert dong["LSX-B"]["thieu"] == pytest.approx(500)


def test_dat_du_thi_phan_du_la_hang_chung(db, svc, customer):
    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500,
              han=HOM_NAY + timedelta(days=20))
    b = _lenh(db, customer, ma="LSX-B", giay_id=g.id, so_to_nguyen=500,
              han=HOM_NAY + timedelta(days=5))
    _, ln = _yc(db, _hang(g), [(a, 500)])
    _don_tu(db, ln, 800)

    dong = _dong(svc.can_doi(), g)
    assert dong["LSX-A"]["thieu"] == pytest.approx(0)
    assert dong["LSX-B"]["thieu"] == pytest.approx(200), "300 dư của đơn là hàng chung"


def test_mot_yeu_cau_hai_lenh_chia_theo_so_de_nghi(db, svc, customer):
    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500, han=MAI)
    b = _lenh(db, customer, ma="LSX-B", giay_id=g.id, so_to_nguyen=300, han=MAI)
    _, ln = _yc(db, _hang(g), [(a, 500), (b, 300)])
    p, pl = _don_tu(db, ln, 800)

    dong = _dong(svc.can_doi(), g)
    assert dong["LSX-A"]["thieu"] == pytest.approx(0)
    assert dong["LSX-B"]["thieu"] == pytest.approx(0)
    assert svc.mach["peg"][pl.id] == {
        (("giay", g.id) + KHO, a.id, None): pytest.approx(500),
        (("giay", g.id) + KHO, b.id, None): pytest.approx(300),
    }


# --- Giữ chỗ lấy phần đặt cho mình trước -------------------------------------------


def test_giu_cho_lenh_khac_khong_giu_duoc_hang_dat_cho_A(db, svc, customer):
    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500,
              han=HOM_NAY + timedelta(days=20))
    b = _lenh(db, customer, ma="LSX-B", giay_id=g.id, so_to_nguyen=500,
              han=HOM_NAY + timedelta(days=5))
    _, ln = _yc(db, _hang(g), [(a, 500)])
    _, pl = _don_tu(db, ln, 500)
    giu = GiuChoService(db, svc)

    assert giu.bat(lsx_id=b.id)["du"] is False, "B không được bám lô đặt cho A"
    assert giu.bat(lsx_id=a.id)["du"] is True
    rows = giu.repo.cua_chu_the(lsx_id=a.id, bai_ghep_id=None)
    assert any(r.nguon == NGUON_DANG_VE and r.purchase_request_line_id == pl.id for r in rows)


# --- Lưu form: gộp liên kết theo ô + chặn ô đã có phiếu -------------------------------


def test_lien_ket_hai_buoc_cung_o_gop_mot_va_cong_so():
    lines = [type("L", (), {"hang_loai": "giay", "hang_id": 7})()]
    ra = PurchaseService._clean_nguon_lenh([
        {"hang_loai": "giay", "hang_id": 7, "kho_rong": 780, "kho_dai": 905, "lsx_id": 1,
         "buoc_id": 1, "so_luong": 100},
        {"hang_loai": "giay", "hang_id": 7, "kho_rong": 780, "kho_dai": 905, "lsx_id": 1,
         "buoc_id": 2, "so_luong": 50},
        {"hang_loai": "giay", "hang_id": 7, "kho_rong": 780, "kho_dai": 905, "lsx_id": 2,
         "buoc_id": 3},
    ], lines)
    assert [(n["lsx_id"], n["kho_rong"], n["so_luong"]) for n in ra] == [
        (1, 780, 150), (2, 780, None)]


def test_luu_yeu_cau_cho_o_da_co_phieu_bi_TU_CHOI(client, db, customer):
    g = _giay(db)
    s = Supplier(name="NCC giấy", status="active")
    db.add(s)
    db.flush()
    db.add(SupplierItem(supplier_id=s.id, hang_loai="giay", hang_id=g.id,
                        item_name=g.ten, unit=_dv_dong(_hang(g), None), unit_price=1))
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500, han=MAI)
    _yc(db, _hang(g), [(a, 500)])
    payload = {
        "source_type": "san_xuat", "purpose": "Thiếu giấy",
        "needed_date": MAI.isoformat(),
        "lines": [{"item_name": g.ten, "unit": _dv_dong(_hang(g), None), "quantity": 500,
                   "hang_loai": "giay", "hang_id": g.id,
                   "kho_rong": KHO[0], "kho_dai": KHO[1]}],
        "nguon_lenh": [{"hang_loai": "giay", "hang_id": g.id, "kho_rong": KHO[0],
                        "kho_dai": KHO[1], "lsx_id": a.id, "so_luong": 500}],
    }
    headers = {"Authorization": f"Bearer {_admin_token()}"}
    resp = client.post("/api/department-purchase-requests", json=payload, headers=headers)
    assert resp.status_code == 409, resp.text
    assert "LSX-A" in resp.json()["detail"] and "YCMH-1" in resp.json()["detail"]
    assert db.query(DepartmentPurchaseRequest).count() == 1


# --- Lưới ---------------------------------------------------------------------------


def _luoi(db, svc):
    from app.services.luoi_vat_tu import dung_luoi

    return dung_luoi(GiuChoService(db, svc))


def _o(luoi, lenh, g):
    return next(d for d in luoi["dong"] if d["lsx_id"] == lenh.id and d["hang_id"] == g.id)


def test_luoi_tinh_trang_tung_o(db, svc, customer):
    from app.services.giu_cho_service import GiuChoService as _G

    g = _giay(db)
    from tests.test_ke_hoach_vat_tu import _ton

    _ton(db, g, 300)
    co = _lenh(db, customer, ma="LSX-CO", giay_id=g.id, so_to_nguyen=300,
               han=HOM_NAY + timedelta(days=1))
    mua = _lenh(db, customer, ma="LSX-MUA", giay_id=g.id, so_to_nguyen=200,
                han=HOM_NAY + timedelta(days=2))
    dang = _lenh(db, customer, ma="LSX-DANG", giay_id=g.id, so_to_nguyen=100,
                 han=HOM_NAY + timedelta(days=3))
    ve = _lenh(db, customer, ma="LSX-VE", giay_id=g.id, so_to_nguyen=400,
               han=HOM_NAY + timedelta(days=4))
    _yc(db, _hang(g), [(dang, 100)])
    _, ln = _yc(db, _hang(g), [(ve, 400)])
    _don_tu(db, ln, 400, ngay_ve=MAI + timedelta(days=5))

    luoi = _luoi(db, svc)
    assert _o(luoi, co, g)["tinh_trang"] == "chua_giu", "tồn đủ mà chưa bật giữ"
    o_mua = _o(luoi, mua, g)
    assert o_mua["tinh_trang"] == "can_mua" and o_mua["khoa_mua"] and o_mua["so_de_nghi"] == 200
    o_dang = _o(luoi, dang, g)
    assert o_dang["tinh_trang"] == "dang_mua" and o_dang["phieu"]["ma"] == "YCMH-1"
    assert not o_dang["khoa_mua"], "đã có phiếu thì không còn khoá mua"
    assert _o(luoi, ve, g)["tinh_trang"] == "chua_giu", "hàng đặt cho lệnh đang về, chưa giữ"

    giu = _G(db, svc)
    giu.bat(lsx_id=co.id)
    giu.bat(lsx_id=ve.id)
    luoi = _luoi(db, svc)
    o_co = _o(luoi, co, g)
    assert o_co["tinh_trang"] == "co_kho" and o_co["ngay_co_hang"]["loai"] == "co_san"
    o_ve = _o(luoi, ve, g)
    assert o_ve["tinh_trang"] == "cho_ve"
    assert o_ve["ngay_co_hang"] == {"loai": "ngay_ve", "ngay": (MAI + timedelta(days=5)).isoformat()}
    assert o_ve["phieu"]["buoc"] == BUOC_DAT and o_ve["da_giu"] == 400


def test_luoi_ghi_chu_trung_va_loc_dem(db, svc, customer):
    from app.services.luoi_vat_tu import chon_trang

    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500, han=MAI)
    b = _lenh(db, customer, ma="LSX-B", giay_id=g.id, so_to_nguyen=100, han=MAI)
    _yc(db, _hang(g), [(a, 500)])
    _yc(db, _hang(g), [(a, 500)])

    luoi = _luoi(db, svc)
    o_a = _o(luoi, a, g)
    assert "1 phiếu trùng" in o_a["ghi_chu"]
    nen_huy = luoi["hang"][o_a["hang"]]["nen_huy"]
    assert [(n["yc_ma"], n["ly_do"]) for n in nen_huy] == [("YCMH-2", "lệnh đã có phiếu YCMH-1")]
    assert nen_huy[0]["co_the_huy"] is True

    trang = chon_trang(luoi, xem="lenh", tinh_trang="can_mua")
    assert trang["dem"]["can_mua"] == 1 and trang["dem"]["dang_mua"] == 1
    assert [n["lenh"]["ma"] for n in trang["items"]] == ["LSX-B"]
    theo_hang = chon_trang(luoi, xem="hang")
    assert len(theo_hang["items"]) == 1 and len(theo_hang["items"][0]["dong"]) == 2
    assert chon_trang(luoi, q="ycmh-2")["tong_dong"] == 1, "tìm theo mã phiếu trùng"


def test_api_luoi_va_hai_ngan(client, db, customer):
    g = _giay(db)
    a = _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500, han=MAI)
    headers = {"Authorization": f"Bearer {_admin_token()}"}
    r = client.get("/api/ke-hoach-vat-tu/luoi", headers=headers)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["items"][0]["lenh"]["ma"] == "LSX-A"
    assert body["items"][0]["dong"][0]["tinh_trang"] == "can_mua"
    r = client.get(f"/api/ke-hoach-vat-tu/luoi/lenh?lsx_id={a.id}", headers=headers)
    assert r.status_code == 200, r.text
    assert r.json()["dong"][0]["buoc"][0]["ten_viec"] == "In offset"
    r = client.get(f"/api/ke-hoach-vat-tu/luoi/hang?hang_loai=giay&hang_id={g.id}"
                   f"&kho_rong={KHO[0]}&kho_dai={KHO[1]}", headers=headers)
    assert r.status_code == 200, r.text
    assert r.json()["hang"]["con_thieu"] == 500


def test_api_luoi_xuat_excel(client, db, customer):
    from io import BytesIO

    from openpyxl import load_workbook

    g = _giay(db)
    _lenh(db, customer, ma="LSX-A", giay_id=g.id, so_to_nguyen=500, han=MAI)
    headers = {"Authorization": f"Bearer {_admin_token()}"}
    r = client.get("/api/ke-hoach-vat-tu/luoi/xuat.xlsx", headers=headers)
    assert r.status_code == 200, r.text
    ws = load_workbook(BytesIO(r.content)).active
    assert ws["A2"].value == "LSX-A" and ws["N2"].value == "Cần mua" and ws["J2"].value == 500


def test_yeu_cau_khong_gan_lenh_chi_goi_y_huy_khi_la_cho_lenh(db, svc, customer):
    """08/10/2026: Mua tồn / Theo yêu cầu cố ý không gắn lệnh — lệnh hết thiếu không phải lý do huỷ.
    Chỉ yêu cầu Cho lệnh SX mất liên kết mới vào danh sách nên huỷ."""
    from app.models.purchase import LOAI_MUA_CHO_LSX, LOAI_MUA_THEO_YEU_CAU, LOAI_MUA_TON
    from tests.test_ke_hoach_vat_tu import _ton

    g = _giay(db)
    _ton(db, g, 300)
    a = _lenh(db, customer, ma="LSX-DU", giay_id=g.id, so_to_nguyen=100, han=MAI)
    GiuChoService(db, svc).bat(lsx_id=a.id)  # giữ đủ từ tồn ⇒ không lệnh nào còn thiếu
    for loai in (LOAI_MUA_TON, LOAI_MUA_THEO_YEU_CAU, LOAI_MUA_CHO_LSX):
        yc, _ = _yc(db, _hang(g), [], so_luong=50)
        yc.loai_mua = loai
    db.commit()

    luoi = _luoi(db, svc)
    nen_huy = luoi["hang"][_o(luoi, a, g)["hang"]]["nen_huy"]
    assert [n["yc_ma"] for n in nen_huy] == ["YCMH-3"]
