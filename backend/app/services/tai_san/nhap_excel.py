"""Nhập TÀI SẢN ĐANG DÙNG từ .xlsx — máy đã chạy trước khi lên phần mềm (làm lại 05/10/2026).

Ngày đầu đưa phần mềm vào dùng là lúc phải nhập dồn vài chục món, nên có đường file. Mỗi dòng đi
qua đúng `TaiSanService.nap_dau_ky` như form "Thêm tài sản đang dùng" — không có luật thứ hai.

Cả file là MỘT giao dịch: một dòng sai là không dòng nào được ghi. `ghi=False` (xem trước) chạy y hệt
rồi rollback, nên con số xem trước là con số thật.

Ô trống được tự điền như form:
- Loại trống ⇒ theo giá một cái: từ 30 triệu là Tài sản cố định, dưới là Công cụ dụng cụ.
- Số lượng trống ⇒ 1. Tài sản cố định luôn một cái một dòng.
- Tính tiếp trên phần mềm từ tháng trống ⇒ tháng hiện tại.
- Đã khấu hao mấy tháng trống ⇒ số tháng từ Bắt đầu dùng tới tháng tính tiếp (tối đa bằng số tháng).
- Đã khấu hao trước đó trống ⇒ Giá mua × số tháng đã khấu hao ÷ Khấu hao trong (làm tròn xuống).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

from sqlalchemy import select
from sqlalchemy.orm import Session

from ...models.department import Department
from ...models.tai_san import LOAI_CCDC, LOAI_TSCD, NGUON_DAU_KY
from .service import TaiSanService, TaiSanTrung, TaiSanValidationError

#: Từ ngưỡng này (một cái) luật tính là tài sản cố định — TT 45/2013 Điều 3.
NGUONG_TSCD = 30_000_000

SHEET = "Tài sản đang dùng"

#: (khoá, tiêu đề trên file, bắt buộc). Tiêu đề khớp không phân biệt hoa thường và bỏ qua dấu `*`.
COT = [
    ("ten", "Tên tài sản", True),
    ("loai", "Loại", False),
    ("gia", "Giá mua một cái (đồng)", True),
    ("so_luong", "Số lượng", False),
    ("so_thang", "Khấu hao trong (tháng)", True),
    ("ngay_su_dung", "Bắt đầu dùng từ ngày", True),
    ("thang_da", "Đã khấu hao mấy tháng", False),
    ("hao_mon", "Đã khấu hao trước đó (đồng)", False),
    ("bo_phan", "Bộ phận dùng", False),
    ("tinh_tu", "Tính tiếp trên phần mềm từ tháng", False),
]

HUONG_DAN = [
    ("Tên tài sản", "Bắt buộc. Ví dụ: Máy in Komori 4 màu."),
    ("Loại", "Tài sản cố định hoặc Công cụ dụng cụ. Bỏ trống: từ 30 triệu một cái là Tài sản cố "
             "định, dưới 30 triệu là Công cụ dụng cụ."),
    ("Giá mua một cái (đồng)", "Bắt buộc. Cộng cả vận chuyển, lắp đặt, chạy thử."),
    ("Số lượng", "Bỏ trống là 1. Tài sản cố định mỗi cái một dòng."),
    ("Khấu hao trong (tháng)", "Bắt buộc. Máy ngành in thường 84–180 tháng. Công cụ dụng cụ tối "
                               "đa 36 tháng."),
    ("Bắt đầu dùng từ ngày", "Bắt buộc. Ngày máy bắt đầu chạy, dạng 15/03/2021."),
    ("Đã khấu hao mấy tháng", "Bỏ trống: phần mềm tự đếm từ ngày bắt đầu dùng tới tháng tính tiếp."),
    ("Đã khấu hao trước đó (đồng)", "Có sổ cũ thì gõ đúng số trên sổ. Bỏ trống: phần mềm tự chia "
                                    "đều theo số tháng."),
    ("Bộ phận dùng", "Gõ đúng tên bộ phận như trên phần mềm. Bỏ trống được."),
    ("Tính tiếp trên phần mềm từ tháng", "Dạng 10/2026. Bỏ trống là tháng này."),
]

_TEN_LOAI = {
    "tài sản cố định": LOAI_TSCD, "tscđ": LOAI_TSCD, "tscd": LOAI_TSCD,
    "công cụ dụng cụ": LOAI_CCDC, "ccdc": LOAI_CCDC,
}


class ExcelSaiMan(Exception):
    """File không đọc được hoặc không đúng mẫu → 422 ở router."""


@dataclass
class KetQuaNhap:
    tong_dong: int = 0
    tao_moi: int = 0
    da_ghi: bool = False
    #: (số dòng trên file, cột, lý do)
    loi: list[tuple[int, str, str]] = field(default_factory=list)


class _LoiO(Exception):
    def __init__(self, cot: str, ly_do: str) -> None:
        super().__init__(ly_do)
        self.cot = cot
        self.ly_do = ly_do


def tao_mau() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    for i, (_, tieu_de, bat_buoc) in enumerate(COT, start=1):
        o = ws.cell(row=1, column=i, value=f"{tieu_de} *" if bat_buoc else tieu_de)
        o.font = Font(bold=True)
        ws.column_dimensions[o.column_letter].width = max(16, len(tieu_de) + 4)
    hd = wb.create_sheet("Hướng dẫn")
    hd.append(["Cột", "Cách điền"])
    hd["A1"].font = hd["B1"].font = Font(bold=True)
    for dong in HUONG_DAN:
        hd.append(list(dong))
    hd.column_dimensions["A"].width = 34
    hd.column_dimensions["B"].width = 100
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _chuan(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").replace("*", "")).strip().lower()


def _so(v, cot: str) -> int | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        if v < 0:
            raise _LoiO(cot, "Không được âm")
        return int(round(v))
    chu = re.sub(r"[^\d]", "", str(v))
    if not chu:
        raise _LoiO(cot, f"'{v}' không phải số")
    return int(chu)


def _ngay(v, cot: str) -> date | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.fullmatch(r"\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})\s*", str(v))
    if m:
        try:
            return date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            pass
    raise _LoiO(cot, f"'{v}' không phải ngày (dạng 15/03/2021)")


def _thang(v, cot: str) -> date | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (datetime, date)):
        return date(v.year, v.month, 1)
    m = re.fullmatch(r"\s*(\d{1,2})[/.-](\d{4})\s*", str(v))
    if m and 1 <= int(m[1]) <= 12:
        return date(int(m[2]), int(m[1]), 1)
    raise _LoiO(cot, f"'{v}' không phải tháng (dạng 10/2026)")


def _so_thang_giua(tu: date, den: date) -> int:
    return (den.year * 12 + den.month) - (tu.year * 12 + tu.month)


def _dong_thanh_payload(o: dict, bo_phan: dict[str, int], hom_nay: date) -> dict:
    tieu = dict((k, t) for k, t, _ in COT)
    ten = str(o.get("ten") or "").strip()
    if not ten:
        raise _LoiO(tieu["ten"], "Phải có tên tài sản")
    gia = _so(o.get("gia"), tieu["gia"])
    if not gia:
        raise _LoiO(tieu["gia"], "Phải có giá mua lớn hơn 0")
    so_luong = _so(o.get("so_luong"), tieu["so_luong"]) or 1
    so_thang = _so(o.get("so_thang"), tieu["so_thang"])
    if not so_thang:
        raise _LoiO(tieu["so_thang"], "Phải có số tháng khấu hao lớn hơn 0")
    nsd = _ngay(o.get("ngay_su_dung"), tieu["ngay_su_dung"])
    if nsd is None:
        raise _LoiO(tieu["ngay_su_dung"], "Phải có ngày bắt đầu dùng")

    chu_loai = _chuan(o.get("loai"))
    if chu_loai:
        loai = _TEN_LOAI.get(chu_loai)
        if loai is None:
            raise _LoiO(tieu["loai"], "Ghi Tài sản cố định hoặc Công cụ dụng cụ")
    else:
        loai = LOAI_TSCD if gia >= NGUONG_TSCD else LOAI_CCDC
    if loai == LOAI_TSCD and so_luong != 1:
        raise _LoiO(tieu["so_luong"], "Tài sản cố định mỗi cái một dòng (số lượng 1)")

    tinh_tu = _thang(o.get("tinh_tu"), tieu["tinh_tu"]) or date(hom_nay.year, hom_nay.month, 1)
    if (nsd.year, nsd.month) > (tinh_tu.year, tinh_tu.month):
        raise _LoiO(tieu["ngay_su_dung"],
                    "Bắt đầu dùng sau tháng tính tiếp — máy mới mua thì dùng nút Thêm tài sản")
    tong = gia * so_luong
    thang_da = _so(o.get("thang_da"), tieu["thang_da"])
    if thang_da is None:
        thang_da = min(max(_so_thang_giua(nsd, tinh_tu), 0), so_thang)
    hao_mon = _so(o.get("hao_mon"), tieu["hao_mon"])
    if hao_mon is None:
        hao_mon = tong if thang_da >= so_thang else tong * thang_da // so_thang

    bo_phan_id = None
    ten_bp = str(o.get("bo_phan") or "").strip()
    if ten_bp:
        bo_phan_id = bo_phan.get(ten_bp.lower())
        if bo_phan_id is None:
            raise _LoiO(tieu["bo_phan"], f"Không có bộ phận tên '{ten_bp}'")

    return {
        "ten": ten, "loai": loai, "so_luong": so_luong, "don_gia": gia,
        "so_thang": so_thang, "ngay_su_dung": nsd, "nguon_vao": NGUON_DAU_KY,
        "moc_tu_ngay": tinh_tu, "hao_mon_dau_ky": hao_mon, "thang_da_trich_dau_ky": thang_da,
        "bo_phan_id": bo_phan_id,
    }


def nhap_dang_dung(
    db: Session,
    svc: TaiSanService,
    du_lieu: bytes,
    *,
    user_id: int | None,
    ghi: bool,
    hom_nay: date | None = None,
) -> KetQuaNhap:
    from openpyxl import load_workbook

    if not du_lieu:
        raise ExcelSaiMan("File rỗng.")
    try:
        wb = load_workbook(BytesIO(du_lieu), read_only=True, data_only=True)
    except Exception:
        raise ExcelSaiMan("Không đọc được file — phải là .xlsx theo file mẫu.") from None
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    dong_iter = ws.iter_rows(values_only=True)
    tieu_de = next(dong_iter, None) or ()
    vi_tri: dict[str, int] = {}
    theo_ten = {_chuan(t): k for k, t, _ in COT}
    for i, ten_cot in enumerate(tieu_de):
        k = theo_ten.get(_chuan(ten_cot))
        if k:
            vi_tri[k] = i
    thieu = [t for k, t, bb in COT if bb and k not in vi_tri]
    if thieu:
        raise ExcelSaiMan(
            "File thiếu cột " + ", ".join(thieu) + ". Bấm Tải file mẫu rồi điền vào đó."
        )

    bo_phan = {
        str(ten).strip().lower(): i
        for i, ten in db.execute(select(Department.id, Department.name))
    }
    hom_nay = hom_nay or date.today()
    kq = KetQuaNhap()
    for so_dong, dong in enumerate(dong_iter, start=2):
        o = {k: (dong[i] if i < len(dong) else None) for k, i in vi_tri.items()}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in o.values()):
            continue
        kq.tong_dong += 1
        try:
            payload = _dong_thanh_payload(o, bo_phan, hom_nay)
            svc.nap_dau_ky(payload, user_id=user_id, commit=False)
            kq.tao_moi += 1
        except _LoiO as e:
            kq.loi.append((so_dong, e.cot, e.ly_do))
        except (TaiSanValidationError, TaiSanTrung) as e:
            kq.loi.append((so_dong, "", str(e)))
    if not kq.tong_dong:
        db.rollback()
        raise ExcelSaiMan("File chưa có dòng tài sản nào.")
    if ghi and not kq.loi:
        db.commit()
        kq.da_ghi = True
    else:
        db.rollback()
    return kq
