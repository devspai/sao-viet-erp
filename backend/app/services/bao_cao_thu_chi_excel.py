"""Xuất BÁO CÁO THU CHI ra .xlsx — đúng bảng đang xem trên màn (kỳ, theo mục/người, Lọc, thẻ).

    dòng 1-2  tiêu đề + kỳ
    dòng 4    tiêu đề cột: Mục (hoặc Người nhận, nộp), các cột kỳ, Cộng kỳ
    dòng 5    Chênh lệch thu − chi
    rồi nhóm Thu (dòng cộng nhóm, các mục), nhóm Chi (dòng cộng nhóm, các mục)
"""
from __future__ import annotations

from io import BytesIO

MEDIA_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DINH_DANG_TIEN = "#,##0"


def _dmy(d) -> str:
    return f"{d.day:02d}/{d.month:02d}/{d.year}"


def ten_file(bc: dict) -> str:
    tu, den = bc["tu_ngay"], bc["den_ngay"]
    return f"bao-cao-thu-chi_{tu:%Y%m%d}-{den:%Y%m%d}.xlsx"


def xuat_xlsx(bc: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Thu chi"
    ws["A1"] = "BÁO CÁO THU CHI"
    ws["A1"].font = Font(size=13)
    ws["A2"] = f"Kỳ {_dmy(bc['tu_ngay'])} – {_dmy(bc['den_ngay'])}. Đơn vị: đồng."
    cot = bc["cot"]
    tieu_de = ["Mục" if bc["theo"] == "muc" else "Người nhận, nộp"] + [c["nhan"] for c in cot] \
        + ["Cộng kỳ"]
    for j, t in enumerate(tieu_de, start=1):
        ws.cell(row=4, column=j, value=t).font = Font(color="555555")
    nen = PatternFill("solid", fgColor="F3F5F8")

    def _dong(r: int, ten: str, o: dict, cong: int, to_nen: bool = False) -> None:
        ws.cell(row=r, column=1, value=ten)
        for j, c in enumerate(cot, start=2):
            ws.cell(row=r, column=j, value=o.get(c["khoa"], 0))
        ws.cell(row=r, column=len(cot) + 2, value=cong)
        for j in range(2, len(cot) + 3):
            ws.cell(row=r, column=j).number_format = DINH_DANG_TIEN
        if to_nen:
            for j in range(1, len(cot) + 3):
                ws.cell(row=r, column=j).fill = nen

    r = 5
    _dong(r, "Chênh lệch thu − chi", bc["chenh_lech"]["o"], bc["chenh_lech"]["cong"], True)
    r += 1
    for ben, nhan in (("thu", "Thu"), ("chi", "Chi")):
        rr = [d for d in bc["dong"] if d["ben"] == ben]
        if not rr:
            continue
        n = bc["nhom"][ben]
        _dong(r, nhan, n["o"], n["cong"], True)
        r += 1
        for d in rr:
            _dong(r, "    " + d["ten"], d["o"], d["cong"])
            r += 1
    ws.column_dimensions["A"].width = 40
    for j in range(2, len(cot) + 3):
        ws.column_dimensions[get_column_letter(j)].width = 18
    ws.freeze_panes = "B5"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
